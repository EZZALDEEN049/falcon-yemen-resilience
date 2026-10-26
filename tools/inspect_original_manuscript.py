from docx import Document
from pathlib import Path
p=Path("paper/Manuscript_2026-10-24_Evidence_Status_Updated.docx")
d=Document(p)
print("=== MANUSCRIPT AUDIT ===")
print("paragraphs",len(d.paragraphs),"tables",len(d.tables),"sections",len(d.sections),"inline_shapes",len(d.inline_shapes))
for i,t in enumerate(d.tables):
 print("TABLE",i,"rows",len(t.rows),"cols",len(t.columns))
 for row in t.rows[:3]:print("  ",[c.text[:120] for c in row.cells])
for i,p in enumerate(d.paragraphs):
 s=p.text.strip()
 if not s:continue
 if p.style.name.startswith("Heading") or i<100 or i>len(d.paragraphs)-140 or any(w.lower() in s.lower() for w in ["abstract","Data availability","conclusion","2023","references","IJCIP","2020","doi:"]):
  print(f"[{i:04d}][{p.style.name}] {s[:460]}")
